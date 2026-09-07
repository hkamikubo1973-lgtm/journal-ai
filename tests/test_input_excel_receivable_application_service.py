import copy
from datetime import datetime
import inspect
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook


SRC_DIR = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC_DIR))

import api.receivable as receivable_api  # noqa: E402
from api.journal import app  # noqa: E402
from columns import EPSON_COLUMNS  # noqa: E402
import input_excel_application_service as application  # noqa: E402
from input_excel_service import InputExcelValidationError  # noqa: E402
from journal_registration_service import (  # noqa: E402
    EDIT_FORM_FIELDS,
    build_registration_id,
)
import receivable_persistence_service as persistence  # noqa: E402
from receivable_registration_handoff_service import (  # noqa: E402
    build_settlement_row_id,
)


FIXED_DATETIME = datetime(2026, 9, 7, 12, 34)


class InputExcelReceivableApplicationServiceTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        self.receivables_directory = self.directory / "receivables"
        self.receivables_directory.mkdir()
        self.current_path = self.receivables_directory / "current.csv"
        self.history_path = (
            self.receivables_directory / "receivable_history.csv"
        )
        self.transactions_path = self.directory / "transactions.csv"
        self.current_path.write_bytes(b"current-fixture\r\n")
        self.history_path.write_bytes(b"history-fixture\r\n")
        self.transactions_path.write_bytes(b"transactions-fixture\r\n")
        self.masters = {
            "accounts": [
                {"code": "100", "name": "普通預金"},
                {"code": "101", "name": "当座預金"},
                {"code": "200", "name": "売掛金"},
            ],
            "departments": [
                {"code": "10", "name": "営業部"},
            ],
            "sub_account_relations": [
                {
                    "account_code": "200",
                    "sub_code": "01",
                    "sub_name": "A商事",
                },
            ],
        }
        app.dependency_overrides[
            receivable_api.get_receivables_directory
        ] = lambda: self.receivables_directory
        app.dependency_overrides[
            receivable_api.get_receivable_account_master_snapshot
        ] = lambda: self.masters
        self.client = TestClient(app)

    def tearDown(self):
        app.dependency_overrides.clear()
        self.temporary_directory.cleanup()

    @staticmethod
    def row(**overrides):
        result = {
            "借方科目": "普通預金",
            "貸方科目": "売掛金",
            "貸方補助": "A商事",
            "部門": "営業部",
            "金額": 1000,
            "摘要": "receipt summary",
        }
        result.update(overrides)
        return result

    def save_receipt(self, rows=None, settlement_id="settlement-001"):
        rows = copy.deepcopy(rows or [self.row()])
        receipt_ref = persistence.calculate_idempotency_key_hash(
            f"receipt-{settlement_id}"
        )
        settlement = {
            "settlement_id": settlement_id,
            "settlement_date": "2026-09-07",
            "customer_name": "A商事",
            "payment_amount": 1000,
            "target_total": 1000,
            "difference": 0,
            "source_candidates": [{
                "コード": "C001",
                "未収ID": "R001",
                "請求日": "2026-08-01",
                "請求額": 1000,
                "残高": 1000,
                "消込予定": 1000,
                "未収科目": "売掛金",
                "未収補助": "A商事",
                "部門": "営業部",
                "取引先": "A商事",
                "摘要": "8月分",
            }],
            "rows": rows,
            "created_at": "2026-09-07T12:00:00+00:00",
        }
        receipt = {
            "schema_version": 1,
            "idempotency_key_hash": receipt_ref,
            "request_hash": "a" * 64,
            "transaction_id": settlement_id,
            "settlement_id": settlement_id,
            "settlement": settlement,
            "current_after_hash": "b" * 64,
            "history_after_hash": "c" * 64,
            "committed_at": "2026-09-07T12:00:01+00:00",
        }
        path = persistence.resolve_settlement_receipt_path(
            self.receivables_directory,
            receipt_ref,
        )
        persistence.save_settlement_receipt(path, receipt)
        return settlement, receipt_ref, path

    @staticmethod
    def receivable_item(settlement, receipt_ref, row_index):
        settlement_id = settlement["settlement_id"]
        return {
            "source_type": "receivable_settlement",
            "provenance": {
                "settlement_id": settlement_id,
                "receipt_ref": receipt_ref,
                "row_index": row_index,
                "row_count": len(settlement["rows"]),
                "settlement_row_id": build_settlement_row_id(
                    receipt_ref,
                    settlement_id,
                    row_index,
                ),
            },
        }

    @staticmethod
    def searched_item(marker="searched"):
        values = {
            "voucher_date": "20260907",
            "voucher_no": f"CERT-{marker}",
            "voucher_summary": f"voucher {marker}",
            "debit_account_code": "100",
            "debit_account_name": "現金",
            "debit_sub_code": "",
            "debit_sub_name": "",
            "debit_dept_code": "",
            "debit_dept_name": "",
            "credit_account_code": "200",
            "credit_account_name": "売上",
            "credit_sub_code": "",
            "credit_sub_name": "",
            "credit_dept_code": "",
            "credit_dept_name": "",
            "amount": 500,
            "summary": marker,
            "source_debit_amount": "500",
            "source_credit_amount": "500",
        }
        prepared = {field: values[field] for field in EDIT_FORM_FIELDS}
        base = {column: "" for column in EPSON_COLUMNS}
        return {
            "registration_id": build_registration_id(prepared, base),
            "prepared_journal": prepared,
            "epson_base_row": base,
            "print_metadata": {"print_category": "searched"},
            "print_warnings": [],
        }

    def export(self, items):
        return application.export_input_excel_application_result(
            items,
            receivables_directory=self.receivables_directory,
            journal_master_snapshot=self.masters,
            export_datetime=FIXED_DATETIME,
        )

    @staticmethod
    def worksheet(content):
        return load_workbook(io.BytesIO(content)).active

    def test_legacy_searched_journal_shape_remains_compatible(self):
        with patch.object(
            application,
            "build_receivable_registration_handoff_application_result",
        ) as receipt_handoff:
            result = self.export([self.searched_item()])

        self.assertEqual(result.rows[0]["摘要"], "searched")
        self.assertEqual(result.rows[0]["区分"], "searched")
        receipt_handoff.assert_not_called()

    def test_receivable_single_row_uses_rebuilt_prepared_journal(self):
        settlement, receipt_ref, _ = self.save_receipt()

        result = self.export([
            self.receivable_item(settlement, receipt_ref, 0)
        ])

        self.assertEqual(result.rows[0]["借方科目"], "普通預金")
        self.assertEqual(result.rows[0]["貸方科目"], "売掛金")
        self.assertEqual(result.rows[0]["貸方補助"], "A商事")
        self.assertEqual(result.rows[0]["摘要"], "receipt summary")

    def test_receivable_multi_row_preserves_request_and_receipt_order(self):
        settlement, receipt_ref, _ = self.save_receipt([
            self.row(金額=600, 摘要="first"),
            self.row(借方科目="当座預金", 金額=400, 摘要="second"),
        ])
        items = [
            self.receivable_item(settlement, receipt_ref, index)
            for index in (0, 1)
        ]

        result = self.export(items)

        self.assertEqual(
            [row["摘要"] for row in result.rows],
            ["first", "second"],
        )
        self.assertEqual([row["No"] for row in result.rows], [1, 2])

    def test_receivable_request_order_must_match_receipt_order(self):
        settlement, receipt_ref, _ = self.save_receipt([
            self.row(金額=600),
            self.row(金額=400),
        ])
        items = [
            self.receivable_item(settlement, receipt_ref, index)
            for index in (1, 0)
        ]

        with self.assertRaisesRegex(
            InputExcelValidationError,
            "receipt順ではありません",
        ):
            self.export(items)

    def test_settlement_row_id_mismatch_is_blocked(self):
        settlement, receipt_ref, _ = self.save_receipt()
        item = self.receivable_item(settlement, receipt_ref, 0)
        item["provenance"]["settlement_row_id"] = "f" * 64

        with self.assertRaisesRegex(
            InputExcelValidationError,
            "settlement_row_idが一致しません",
        ):
            self.export([item])

    def test_row_index_out_of_range_is_blocked(self):
        settlement, receipt_ref, _ = self.save_receipt()
        item = self.receivable_item(settlement, receipt_ref, 0)
        item["provenance"]["row_index"] = 1
        item["provenance"]["settlement_row_id"] = build_settlement_row_id(
            receipt_ref,
            settlement["settlement_id"],
            1,
        )

        with self.assertRaisesRegex(InputExcelValidationError, "範囲外"):
            self.export([item])

    def test_row_count_mismatch_is_blocked(self):
        settlement, receipt_ref, _ = self.save_receipt()
        item = self.receivable_item(settlement, receipt_ref, 0)
        item["provenance"]["row_count"] = 2

        with self.assertRaisesRegex(
            InputExcelValidationError,
            "row_countがreceiptと一致しません",
        ):
            self.export([item])

    def test_settlement_id_mismatch_is_blocked(self):
        settlement, receipt_ref, _ = self.save_receipt()
        item = self.receivable_item(settlement, receipt_ref, 0)
        item["provenance"]["settlement_id"] = "different"

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [item]},
        )

        self.assertEqual(response.status_code, 409)

    def test_malformed_receipt_ref_is_422(self):
        settlement, receipt_ref, _ = self.save_receipt()
        item = self.receivable_item(settlement, receipt_ref, 0)
        item["provenance"]["receipt_ref"] = "INVALID"

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [item]},
        )

        self.assertEqual(response.status_code, 422)

    def test_missing_receipt_is_404(self):
        settlement = {
            "settlement_id": "missing",
            "rows": [self.row()],
        }
        item = self.receivable_item(settlement, "d" * 64, 0)

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [item]},
        )

        self.assertEqual(response.status_code, 404)

    def test_corrupt_receipt_is_503(self):
        receipt_ref = "e" * 64
        path = persistence.resolve_settlement_receipt_path(
            self.receivables_directory,
            receipt_ref,
        )
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("{corrupt", encoding="utf-8")
        settlement = {"settlement_id": "corrupt", "rows": [self.row()]}

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [
                self.receivable_item(settlement, receipt_ref, 0)
            ]},
        )

        self.assertEqual(response.status_code, 503)
        self.assertNotIn(str(path), response.text)

    def test_current_master_resolution_failure_is_422(self):
        settlement, receipt_ref, _ = self.save_receipt()
        self.masters["accounts"] = [
            account for account in self.masters["accounts"]
            if account["name"] != "売掛金"
        ]

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [
                self.receivable_item(settlement, receipt_ref, 0)
            ]},
        )

        self.assertEqual(response.status_code, 422)

    def test_frontend_prepared_journal_and_amount_are_not_trusted(self):
        settlement, receipt_ref, _ = self.save_receipt()
        item = self.receivable_item(settlement, receipt_ref, 0)
        item["prepared_journal"] = {
            "amount": 999999,
            "summary": "tampered",
            "debit_account_name": "tampered",
        }
        item["amount"] = 999999

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [item]},
        )

        self.assertEqual(response.status_code, 200)
        worksheet = self.worksheet(response.content)
        self.assertEqual(worksheet["C2"].value, "普通預金")
        self.assertEqual(worksheet["E2"].value, 1000)
        self.assertEqual(worksheet["I2"].value, "receipt summary")

    def test_invalid_item_blocks_workbook_generation_all_or_none(self):
        settlement, receipt_ref, _ = self.save_receipt([
            self.row(金額=600),
            self.row(金額=400),
        ])
        valid = self.receivable_item(settlement, receipt_ref, 0)
        invalid = self.receivable_item(settlement, receipt_ref, 1)
        invalid["provenance"]["settlement_row_id"] = "0" * 64

        with patch.object(
            application,
            "export_validated_input_excel",
        ) as workbook_builder, self.assertRaises(InputExcelValidationError):
            self.export([valid, invalid])

        workbook_builder.assert_not_called()

    def test_mixed_searched_and_receivable_cart_preserves_input_order(self):
        settlement, receipt_ref, _ = self.save_receipt()
        items = [
            self.searched_item("before"),
            self.receivable_item(settlement, receipt_ref, 0),
            self.searched_item("after"),
        ]

        result = self.export(items)

        self.assertEqual(
            [row["摘要"] for row in result.rows],
            ["before", "receipt summary", "after"],
        )
        self.assertEqual([row["No"] for row in result.rows], [1, 2, 3])

    def test_receivable_download_api_returns_xlsx(self):
        settlement, receipt_ref, _ = self.save_receipt()

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [
                self.receivable_item(settlement, receipt_ref, 0)
            ]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet",
        )
        self.assertIn("input_journal_print_", response.headers[
            "content-disposition"
        ])
        self.assertEqual(self.worksheet(response.content).max_column, 12)

    def test_receivable_save_api_writes_xlsx_without_transactions_update(self):
        settlement, receipt_ref, _ = self.save_receipt()
        export_directory = self.directory / "exports"
        export_directory.mkdir()
        transactions_before = self.transactions_path.read_bytes()

        with patch(
            "input_excel_save_service.load_system_settings",
            return_value={"csv_export_dir": str(export_directory)},
        ):
            response = self.client.post(
                "/api/journal/save-input-excel",
                json={"items": [
                    self.receivable_item(settlement, receipt_ref, 0)
                ]},
            )

        self.assertEqual(response.status_code, 200)
        saved_path = Path(response.json()["saved_path"])
        self.assertTrue(saved_path.is_file())
        self.assertEqual(saved_path.suffix, ".xlsx")
        self.assertEqual(
            self.transactions_path.read_bytes(),
            transactions_before,
        )

    def test_download_keeps_ledger_history_receipt_and_transactions_unchanged(self):
        settlement, receipt_ref, receipt_path = self.save_receipt()
        tracked = (
            self.current_path,
            self.history_path,
            receipt_path,
            self.transactions_path,
        )
        before = {path: path.read_bytes() for path in tracked}

        response = self.client.post(
            "/api/journal/export-input-excel",
            json={"items": [
                self.receivable_item(settlement, receipt_ref, 0)
            ]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual({path: path.read_bytes() for path in tracked}, before)
        self.assertFalse(
            (self.receivables_directory / ".transactions").exists()
        )

    def test_union_openapi_exposes_only_receivable_provenance(self):
        schema = self.client.get("/openapi.json").json()
        schemas = schema["components"]["schemas"]
        receivable = schemas["ReceivableSettlementInputExcelItemRequest"]
        self.assertEqual(
            set(receivable["properties"]),
            {"source_type", "provenance"},
        )
        provenance = schemas["ReceivableInputExcelProvenanceRequest"]
        self.assertFalse(provenance["additionalProperties"])
        self.assertEqual(set(provenance["required"]), {
            "settlement_id", "receipt_ref", "row_index", "row_count",
            "settlement_row_id",
        })

    def test_receivable_branch_has_no_epson_service_dependency(self):
        source = inspect.getsource(application)
        for forbidden in (
            "journal_export_service",
            "epson_export_service",
            "epson_base_row",
            "epson_preview_row",
            "registration_id",
        ):
            self.assertNotIn(forbidden, source)


if __name__ == "__main__":
    unittest.main()
